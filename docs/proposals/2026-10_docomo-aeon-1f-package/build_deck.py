# -*- coding: utf-8 -*-
"""ドコモ中央様向け イオンモール1Fメイン展開 集客コンテンツパッケージ（10ライン構成案）

過去資料（DreamCreation 3Q企画書 / DreamTeam振返り資料）のデザインを踏襲:
メイリオ・赤 #CC0000・大きな章番号・淡いピンクのカード・赤ヘッダーの表・
フッター「NTT DOCOMO Confidential」。

使い方: python3 build_deck.py  →  同じフォルダに .pptx を出力
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
IMG = ROOT / "site" / "images"
OUT = HERE / "【ドコモ中央様】イオンモール1F集客コンテンツパッケージ_10ライン構成案_DreamCreation.pptx"

FONT = "メイリオ"
RED = RGBColor(0xCC, 0x00, 0x00)
RED_D = RGBColor(0x8C, 0x00, 0x00)
INK = RGBColor(0x1A, 0x1A, 0x1A)
INK2 = RGBColor(0x4A, 0x4A, 0x4A)
SUB = RGBColor(0x76, 0x76, 0x76)
LINE = RGBColor(0xE0, 0xE0, 0xE0)
PINK = RGBColor(0xFD, 0xF0, 0xF0)
PINK2 = RGBColor(0xF4, 0xCF, 0xCF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0xE2, 0xEF, 0xD9)
GREEN_D = RGBColor(0x38, 0x76, 0x1D)
BLUE = RGBColor(0xE3, 0xEE, 0xFA)
BLUE_D = RGBColor(0x1F, 0x5F, 0xAD)
YEL = RGBColor(0xFF, 0xF4, 0xD6)
YEL_D = RGBColor(0xA6, 0x6A, 0x00)
GRAY = RGBColor(0xF5, 0xF5, 0xF5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 13


# ---------- helpers ----------
def _font(run, size, color=INK, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, lines, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """lines: str or list of str / (str, dict) — dict may override size,color,bold."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        opts = {}
        if isinstance(ln, tuple):
            ln, opts = ln
        r = p.add_run()
        r.text = ln
        _font(r, opts.get("size", size), opts.get("color", color), opts.get("bold", bold))
    return tb


def box(slide, x, y, w, h, fill=PINK, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def label_in(shape, s, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    _font(r, size, color, bold)


def circle_num(slide, x, y, n, d=Inches(0.42), size=13):
    c = box(slide, x, y, d, d, RED, shape=MSO_SHAPE.OVAL)
    label_in(c, str(n), size)
    return c


def photo(slide, path, x, y, w, h):
    """画像を枠いっぱいにトリミングして貼る。無ければ差し込み予定枠。"""
    if path is None or not Path(path).exists():
        ph = box(slide, x, y, w, h, GRAY, LINE)
        ph.line.dash_style = 4
        label_in(ph, "写真差し込み予定\n（試作品撮影後）", 9, SUB, False)
        return ph
    from PIL import Image
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    box_r, img_r = w / h, iw / ih
    if img_r > box_r:
        crop = (1 - box_r / img_r) / 2
        pic.crop_left = pic.crop_right = crop
    else:
        crop = (1 - img_r / box_r) / 2
        pic.crop_top = pic.crop_bottom = crop
    return pic


def table(slide, x, y, w, rows, col_w, row_h=Inches(0.36), head_size=10, body_size=9.5,
          zebra=True, first_bold=True, aligns=None):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, x, y, w, row_h * nr)
    t = gt.table
    tblPr = gt._element.graphic.graphicData.tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
    for i, cw in enumerate(col_w):
        t.columns[i].width = cw
    for r in range(nr):
        t.rows[r].height = row_h
        for c in range(nc):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = RED
            elif zebra and r % 2 == 0:
                cell.fill.fore_color.rgb = PINK
            else:
                cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            val = rows[r][c]
            parts = val.split("\n")
            for k, part in enumerate(parts):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                if r == 0:
                    p.alignment = PP_ALIGN.CENTER
                elif aligns:
                    p.alignment = aligns[c]
                run = p.add_run()
                run.text = part
                if r == 0:
                    _font(run, head_size, WHITE, True)
                else:
                    _font(run, body_size if k == 0 else body_size - 1.5,
                          INK if k == 0 else SUB, bold=(first_bold and c == 0 and k == 0))
    return t


def header(slide, num, title, lead):
    text(slide, Inches(0.5), Inches(0.32), Inches(1.0), Inches(0.75), num, 34, RED, True)
    text(slide, Inches(1.45), Inches(0.42), Inches(11.4), Inches(0.6), title, 24, INK, True)
    text(slide, Inches(1.45), Inches(1.0), Inches(11.4), Inches(0.45), lead, 11, INK2)


def footer(slide, page, note=""):
    ln = slide.shapes.add_connector(1, Inches(0.5), Inches(7.05), Inches(12.83), Inches(7.05))
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)
    text(slide, Inches(0.5), Inches(7.12), Inches(7.5), Inches(0.25),
         f"{page:02d} / {TOTAL}　{note}", 8, SUB)
    text(slide, Inches(5.2), Inches(7.12), Inches(3.0), Inches(0.25),
         "NTT DOCOMO Confidential", 8, SUB, align=PP_ALIGN.CENTER)
    text(slide, Inches(9.3), Inches(7.12), Inches(3.53), Inches(0.25),
         "DREAM CREATION · OCTOBER 2026", 8, SUB, align=PP_ALIGN.RIGHT)


def callout(slide, y, head, body, h=Inches(0.62)):
    b = box(slide, Inches(0.5), y, Inches(12.33), h, PINK2)
    text(slide, Inches(0.75), y + Inches(0.08), Inches(2.1), h - Inches(0.16), head, 11, RED_D,
         True, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(2.9), y + Inches(0.08), Inches(9.7), h - Inches(0.16), body, 10.5, INK,
         True, anchor=MSO_ANCHOR.MIDDLE)
    return b


def tag(slide, x, y, s, fill=RED, color=WHITE, w=None, size=8.5):
    w = w or Inches(0.13 * len(s) + 0.25)
    t = box(slide, x, y, w, Inches(0.24), fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    t.adjustments[0] = 0.5
    label_in(t, s, size, color)
    return t


# ---------- 01 表紙 ----------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, prs.slide_width, Inches(0.18), RED)
text(s, Inches(0.8), Inches(1.25), Inches(11), Inches(0.4),
     "イオンモール 1F メインイベントスペース展開", 16, RED, True)
text(s, Inches(0.8), Inches(1.75), Inches(11.5), Inches(1.5),
     ["ファミリー集客コンテンツ", "パッケージ ご提案（10ライン構成案）"], 34, INK, True, spacing=1.1)
text(s, Inches(0.8), Inches(3.45), Inches(7), Inches(0.9),
     ["ワークショップ × エア遊具 × ゲーム遊具を「役割分担」で組み合わせ、",
      "視認 → 立寄 → 参加 → 会話 → 商談 までを1ブースで完結させる構成案です。"], 12, INK2)
for i, (p, lab) in enumerate([(IMG / "airizm.jpg", "エアリズム"),
                              (IMG / "punilab-paw.jpg", "スクイーズWS"),
                              (IMG / "carnival-b.jpg", "カーニバル")]):
    x = Inches(8.1) + Inches(1.62) * i
    photo(s, p, x, Inches(3.35), Inches(1.5), Inches(1.5))
    text(s, x, Inches(4.9), Inches(1.5), Inches(0.25), lab, 8, SUB, align=PP_ALIGN.CENTER)
ln = s.shapes.add_connector(1, Inches(0.8), Inches(5.55), Inches(12.5), Inches(5.55))
ln.line.color.rgb = LINE
text(s, Inches(0.8), Inches(5.75), Inches(6), Inches(0.9),
     [("2026年10月", {"size": 12, "color": INK2}),
      ("株式会社DreamCreation 企画書", {"size": 14, "bold": True})], 12)
text(s, Inches(7.3), Inches(5.75), Inches(5.2), Inches(0.9),
     ["提出先：NTTドコモ 中央 ご担当者様", "コンテンツ企画・運営協力：株式会社NextVision"], 10, SUB,
     align=PP_ALIGN.RIGHT)
text(s, Inches(5.2), Inches(7.12), Inches(3.0), Inches(0.25), "NTT DOCOMO Confidential", 8, SUB,
     align=PP_ALIGN.CENTER)

# ---------- 02 本提案の要点 ----------
s = prs.slides.add_slide(BLANK)
header(s, "01", "本提案の要点",
       "イオンモール1Fメインに展開できる集客コンテンツを、10ラインのパッケージとしてご提案します。")
points = [
    ("役割分担型パッケージ",
     "1ブースを「アイキャッチ（エア遊具）」「滞在（ワークショップ）」「キッズケア（ゲーム遊具）」の"
     "3役で構成。集客だけで終わらせず、保護者様との会話・商談時間を生み出す設計です。", "→ P.03"),
    ("10ライン構成 × 4タイプ",
     "会場面積・施設ルールに合わせて選べる4タイプ（大型／スタンダード／季節限定／省スペース）で"
     "計10ラインを編成。どのモールの1Fメインにも当てはめやすい型で揃えます。", "→ P.06 / P.07"),
    ("新ワークショップ「スノードーム」",
     "スクイーズ・めじるしチャームに加え、試験段階のスノードーム作りを冬季限定で試験導入。"
     "2ラインで検証後、反応を見て横展開する段階導入をご提案します。", "→ P.10"),
]
for i, (h, b, ref) in enumerate(points):
    y = Inches(1.7) + Inches(1.38) * i
    box(s, Inches(0.5), y, Inches(12.33), Inches(1.2), PINK, PINK2)
    circle_num(s, Inches(0.75), y + Inches(0.38), i + 1)
    text(s, Inches(1.4), y + Inches(0.18), Inches(3.4), Inches(0.85), h, 15, INK, True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(4.8), y + Inches(0.18), Inches(6.9), Inches(0.85), b, 10.5, INK2,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(11.75), y + Inches(0.18), Inches(0.95), Inches(0.85), ref, 9, RED, True,
         anchor=MSO_ANCHOR.MIDDLE)
callout(s, Inches(5.95), "前提",
        "すべてのラインはリースライン内で完結する配置とし、過度な声掛けに頼らず「お客様から近づいていただける」"
        "体験で集客します。")
footer(s, 2)

# ---------- 03 パッケージの考え方 ----------
s = prs.slides.add_slide(BLANK)
header(s, "02", "パッケージの考え方 ─ 3つの役割で1ブースを構成",
       "コンテンツを単体で置くのではなく、来場者の行動段階ごとに役割を割り当てて組み合わせます。")
roles = [
    ("ROLE 01", "アイキャッチ", "エア遊具 / キッズライセンスパーク",
     ["遠目から「何かやっている」が伝わる大型ビジュアル", "通行客の足を止め、立寄率を上げる",
      "例：エアリズム・カーニバルA/B・2GAME in 1"], RED),
    ("ROLE 02", "滞在・参加", "ワークショップ",
     ["「世界に一つの作品づくり」で着座を生む", "作品完成までの待ち時間＝保護者様とのヒアリング時間",
      "例：スクイーズ・めじるしチャーム・スノードーム"], BLUE_D),
    ("ROLE 03", "キッズケア", "ゲーム遊具（イライラ棒）",
     ["商談中のお子様が飽きずに待てる", "保護者様が落ち着いて話を聞ける環境をつくる",
      "例：電撃チャレンジ（イライラ棒）"], GREEN_D),
]
for i, (en, jp, sub, items, col) in enumerate(roles):
    x = Inches(0.5) + Inches(4.17) * i
    y = Inches(1.7)
    box(s, x, y, Inches(3.95), Inches(3.05), WHITE, LINE)
    box(s, x, y, Inches(3.95), Inches(0.09), col)
    text(s, x + Inches(0.25), y + Inches(0.25), Inches(3.5), Inches(0.25), en, 9, col, True)
    text(s, x + Inches(0.25), y + Inches(0.5), Inches(3.5), Inches(0.45), jp, 18, INK, True)
    text(s, x + Inches(0.25), y + Inches(0.98), Inches(3.5), Inches(0.3), sub, 10, col, True)
    text(s, x + Inches(0.25), y + Inches(1.45), Inches(3.5), Inches(1.5),
         ["▍" + it for it in items], 10, INK2, spacing=1.35)
# funnel
steps = ["視認", "立寄", "参加", "会話", "商談"]
owners = ["アイキャッチ", "アイキャッチ", "ワークショップ", "WS待ち時間", "キッズケア"]
for i, (st, ow) in enumerate(zip(steps, owners)):
    x = Inches(0.5) + Inches(2.47) * i
    c = box(s, x, Inches(5.0), Inches(2.3), Inches(0.62), RED if i < 4 else RED_D,
            shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
    label_in(c, st, 14)
    text(s, x, Inches(5.68), Inches(2.3), Inches(0.3), "主役：" + ow, 9, SUB, align=PP_ALIGN.CENTER)
callout(s, Inches(6.1), "現場からの学び",
        "直近の稼働では、イライラ棒は「集客」より「商談中のお子様対応」で効果を発揮しました。"
        "役割を固定して配置することで、各コンテンツの強みを最大化します。", Inches(0.8))
footer(s, 3)

# ---------- 04 ワークショップ ----------
s = prs.slides.add_slide(BLANK)
header(s, "03", "コンテンツラインアップ ① ワークショップ（滞在・参加）",
       "当社保有の3種のワークショップ。低年齢でも参加できる難易度で、親子の着座時間を生み出します。")
ws = [
    ("スクイーズ ワークショップ", "実績あり", RED, IMG / "punilab-paw.jpg",
     [("所要", "約20分／組"), ("対象", "3歳〜小学生・親子"), ("内容", "動物型・カラーを選び、型から剥がして\nマジックパウダーで仕上げ"),
      ("強み", "現場評価が高く、待ち時間をそのまま\nヒアリング時間に転換できる")]),
    ("めじるしチャーム ワークショップ", "実績あり", RED, None,
     [("所要", "約10〜15分／組（要確認）"), ("対象", "未就学〜小学生・親子"),
      ("内容", "パーツを選んで作る、持ち物の\n“めじるし”になるオリジナルチャーム"),
      ("強み", "回転が早く家族単位で集まりやすい\n（直近稼働で集客の立ち上がりが早い）")]),
    ("スノードーム ワークショップ", "試験導入", YEL_D, None,
     [("所要", "約15〜20分／組（試作検証中）"), ("対象", "小学生〜大人・親子"),
      ("内容", "フィギュアやラメを選んで閉じ込める\n冬らしい作品づくり"),
      ("強み", "クリスマス〜冬の季節感で話題化。\n大人も参加でき滞在時間が長い")]),
]
for i, (name, st, col, img, rows) in enumerate(ws):
    x = Inches(0.5) + Inches(4.17) * i
    y = Inches(1.65)
    box(s, x, y, Inches(3.95), Inches(4.85), WHITE, LINE)
    photo(s, img, x + Inches(0.15), y + Inches(0.15), Inches(3.65), Inches(1.75))
    tag(s, x + Inches(0.15), y + Inches(2.05), st, col, w=Inches(1.0))
    text(s, x + Inches(0.15), y + Inches(2.36), Inches(3.65), Inches(0.4), name, 13, INK, True)
    yy = y + Inches(2.85)
    for k, v in rows:
        text(s, x + Inches(0.15), yy, Inches(0.6), Inches(0.3), k, 9, col, True)
        text(s, x + Inches(0.75), yy, Inches(3.05), Inches(0.5), v.split("\n"), 9, INK2, spacing=1.1)
        yy += Inches(0.47)
text(s, Inches(0.5), Inches(6.6), Inches(12.3), Inches(0.35),
     "※ 参考価格（最新カタログ・税別）：スクイーズWS ¥70,000／日。めじるしチャーム・スノードームは別途お見積り。"
     "材料費・送料は別途。", 9, SUB)
footer(s, 4)

# ---------- 05 エア遊具・ゲーム ----------
s = prs.slides.add_slide(BLANK)
header(s, "04", "コンテンツラインアップ ② エア遊具・ゲーム遊具",
       "NextVisionが自社保有・運用する大型コンテンツ。1Fメインの吹き抜けでも遠目から目に留まります。")
items = [
    ("エアリズム", "アイキャッチ", IMG / "airizm.jpg", "W2800×D1000×H1750", "¥50,000",
     "風船が舞う屋内向けドーム。安全性と視認性を両立"),
    ("カーニバルA", "アイキャッチ", IMG / "carnival-a.jpg", "W1500×D1450×H2000", "¥50,000",
     "短時間回転で行列をつくりやすい的当てゲーム"),
    ("カーニバルB", "アイキャッチ", IMG / "carnival-b.jpg", "W1500×D1450×H2000", "¥50,000",
     "兄弟・家族で同時参加でき、滞在を延ばせる"),
    ("2GAME in 1", "アイキャッチ", HERE / "img" / "2game.png", "W2700×D2000×H3000", "¥50,000",
     "高さ3mのテント型。吹き抜けで最も目立つ"),
    ("電撃チャレンジ（イライラ棒）", "キッズケア", HERE / "img" / "dengeki.png", "W1260×D485×H1460", "¥80,000",
     "商談中のお子様が夢中になれる。省スペース"),
    ("キッズライセンスパーク", "アイキャッチ", IMG / "klp.jpg", "コース約6×4m〜", "¥80,000〜",
     "EVカー運転＋こども免許証。記念性・撮影価値"),
]
for i, (name, role, img, size, price, desc) in enumerate(items):
    col, row = i % 3, i // 3
    x = Inches(0.5) + Inches(4.17) * col
    y = Inches(1.6) + Inches(2.58) * row
    box(s, x, y, Inches(3.95), Inches(2.42), WHITE, LINE)
    photo(s, img, x + Inches(0.12), y + Inches(0.12), Inches(1.6), Inches(2.18))
    tag(s, x + Inches(1.85), y + Inches(0.15), role, GREEN_D if role == "キッズケア" else RED,
        w=Inches(1.05))
    text(s, x + Inches(1.85), y + Inches(0.48), Inches(2.0), Inches(0.6), name, 12, INK, True)
    text(s, x + Inches(1.85), y + Inches(1.08), Inches(2.0), Inches(0.7), desc, 8.5, INK2)
    text(s, x + Inches(1.85), y + Inches(1.72), Inches(2.0), Inches(0.6),
         [("本体 " + size, {"size": 7.5, "color": SUB}),
          (price + " ／日（税別・参考）", {"size": 9, "color": RED, "bold": True})], 8)
footer(s, 5, "価格は最新カタログの参考価格（税別・送料別途）")

# ---------- 06 4タイプ ----------
s = prs.slides.add_slide(BLANK)
header(s, "05", "パッケージ 4タイプ ─ 会場面積と施設ルールで選べる型",
       "1Fメインの広さ・天井高・リースラインの厳しさに応じて、4つの型から選択いただけます。")
types = [
    ("TYPE A", "大型フル", "約10m×8m〜", RED,
     "エア遊具 ＋ ワークショップ ＋ イライラ棒", "吹き抜け・大型催事スペース", "集客力最大。週末・大型連休向け"),
    ("TYPE B", "スタンダード", "約8m×6m〜", BLUE_D,
     "ゲーム遊具 ＋ ワークショップ（＋イライラ棒）", "標準的な1Fイベントスペース", "最も汎用性が高い基本形"),
    ("TYPE C", "季節限定", "約8m×6m〜", YEL_D,
     "エア遊具 or ゲーム ＋ スノードーム", "11月〜2月のクリスマス・冬催事", "季節感で話題化。試験導入枠"),
    ("TYPE D", "省スペース", "約5m×4m〜", GREEN_D,
     "ワークショップ2種 ＋ イライラ棒", "狭小スペース・規制の厳しい施設", "リースライン内で確実に完結"),
]
for i, (en, jp, area, col, comp, scene, merit) in enumerate(types):
    x = Inches(0.5) + Inches(3.12) * i
    y = Inches(1.65)
    box(s, x, y, Inches(2.95), Inches(3.9), WHITE, LINE)
    hd = box(s, x, y, Inches(2.95), Inches(0.95), col)
    text(s, x + Inches(0.2), y + Inches(0.1), Inches(2.6), Inches(0.3), en, 10, WHITE, True)
    text(s, x + Inches(0.2), y + Inches(0.38), Inches(2.6), Inches(0.5), jp, 18, WHITE, True)
    yy = y + Inches(1.15)
    for k, v in (("想定面積", area), ("構成", comp), ("想定会場", scene), ("ねらい", merit)):
        text(s, x + Inches(0.2), yy, Inches(2.6), Inches(0.25), k, 9, col, True)
        text(s, x + Inches(0.2), yy + Inches(0.25), Inches(2.6), Inches(0.45), v, 10, INK)
        yy += Inches(0.68)
callout(s, Inches(5.8), "選定の進め方",
        "ロケハン時に面積・天井高・音出し可否・リースラインを実測し、型を確定します。"
        "規制が厳しい施設では無理に大型を持ち込まず、TYPE Dに切り替えます。", Inches(0.8))
footer(s, 6, "面積は目安。各モールの施設規定・設置図により調整します。")

# ---------- 07 10ライン構成 ----------
s = prs.slides.add_slide(BLANK)
header(s, "06", "10ライン構成案",
       "4タイプを組み合わせ、同時期に10ラインを稼働できる編成です。ラインごとにメインのWSを分散させ、在庫を平準化します。")
rows = [
    ["No.", "ライン名（仮）", "型", "アイキャッチ", "ワークショップ", "キッズケア", "推奨時期"],
    ["L-01", "ふわふわバルーンパーク", "A", "エアリズム", "スクイーズ", "イライラ棒", "通年"],
    ["L-02", "バルーン×チャーム工房", "A", "エアリズム", "めじるしチャーム", "イライラ棒", "通年"],
    ["L-03", "ビッグカーニバル", "A", "2GAME in 1", "スクイーズ", "イライラ棒", "通年・連休"],
    ["L-04", "キッズドライブ＆クラフト", "A", "キッズライセンスパーク", "めじるしチャーム", "─", "通年"],
    ["L-05", "カーニバル工房", "B", "カーニバルA", "スクイーズ", "イライラ棒", "通年"],
    ["L-06", "ファミリーチャレンジ", "B", "カーニバルB", "めじるしチャーム", "イライラ棒", "通年"],
    ["L-07", "スノードーム・ウィンター", "C", "エアリズム", "スノードーム", "イライラ棒", "11月〜2月"],
    ["L-08", "ウィンターカーニバル", "C", "カーニバルB", "スノードーム＋チャーム", "─", "11月〜2月"],
    ["L-09", "クラフトステーション", "D", "─（装飾・POP）", "スクイーズ＋チャーム", "イライラ棒", "通年"],
    ["L-10", "ミニ工房", "D", "─（装飾・POP）", "スクイーズ or チャーム", "イライラ棒", "通年"],
]
C = PP_ALIGN.CENTER
L = PP_ALIGN.LEFT
table(s, Inches(0.5), Inches(1.6), Inches(12.33), rows,
      [Inches(0.75), Inches(2.85), Inches(0.6), Inches(2.4), Inches(2.4), Inches(1.53), Inches(1.8)],
      row_h=Inches(0.4), aligns=[C, L, C, L, L, C, C])
text(s, Inches(0.5), Inches(6.12), Inches(12.33), Inches(0.8),
     ["▍WS内訳：スクイーズ L-01/03/05/09/10　めじるしチャーム L-02/04/06/08/09/10　スノードーム L-07/08",
      "▍アイキャッチ内訳：エアリズム 3／カーニバルA 1／カーニバルB 2／2GAME in 1 1／キッズライセンスパーク 1／TYPE D 2",
      "▍同時稼働時の機材在庫（特にエアリズム3基・イライラ棒8台）は確認のうえ、不足分は追加調達または時期分散で対応します。"],
     9.5, INK2, spacing=1.3)
footer(s, 7, "ライン名は仮称。施設・時期に応じて入れ替え可能です。")

# ---------- 08 レイアウトイメージ ----------
s = prs.slides.add_slide(BLANK)
header(s, "07", "1Fメイン レイアウトイメージ（TYPE A 例）",
       "通路側にアイキャッチ、中央にワークショップ、奥に商談席。お子様の動線と保護者様の動線を分けます。")
# area
ax, ay, aw, ah = Inches(0.6), Inches(1.75), Inches(7.4), Inches(4.9)
area = box(s, ax, ay, aw, ah, RGBColor(0xFA, 0xFA, 0xFA), RED, lw=2)
area.line.dash_style = 7
text(s, ax + Inches(0.1), ay - Inches(0.3), Inches(4), Inches(0.25), "リースライン（施設様と実測・床面表示）",
     9, RED, True)
# aisle
aisle = box(s, ax, ay + ah + Inches(0.05), aw, Inches(0.22), LINE)
label_in(aisle, "← 通路（メイン動線）→", 8, SUB, False)


def zone(x, y, w, h, fill, line, title, sub):
    z = box(s, x, y, w, h, fill, line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + Inches(0.1), y + Inches(0.08), w - Inches(0.2), h - Inches(0.16),
         [(title, {"bold": True, "size": 11}), (sub, {"size": 8.5, "color": INK2})], 11,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


zone(ax + Inches(0.3), ay + Inches(3.0), Inches(3.0), Inches(1.6), PINK, RED, "① アイキャッチ",
     "エアリズム等（通路側・視認）")
zone(ax + Inches(3.6), ay + Inches(3.3), Inches(1.5), Inches(1.3), GREEN, GREEN_D, "受付",
     "お得診断アンケート")
zone(ax + Inches(5.3), ay + Inches(2.4), Inches(1.8), Inches(2.2), BLUE, BLUE_D, "② WS テーブル",
     "6〜8席")
zone(ax + Inches(0.3), ay + Inches(0.3), Inches(4.6), Inches(1.6), WHITE, INK2, "④ 商談テーブル",
     "保護者様のヒアリング・ご提案")
zone(ax + Inches(5.3), ay + Inches(0.3), Inches(1.8), Inches(1.8), GREEN, GREEN_D, "③ イライラ棒",
     "商談席から見える位置")
# flow notes
flow = [("1", "通路から①が目に入り、立ち寄り"), ("2", "受付でお得診断アンケートを記入しWSへ"),
        ("3", "WS制作中に保護者様へヒアリング"), ("4", "④商談テーブルへご案内"),
        ("5", "お子様は③で遊びながら待機")]
text(s, Inches(8.4), Inches(1.75), Inches(4.4), Inches(0.3), "来場者の動き", 13, INK, True)
for i, (n, t) in enumerate(flow):
    y = Inches(2.2) + Inches(0.62) * i
    circle_num(s, Inches(8.4), y, n, Inches(0.38), 11)
    text(s, Inches(8.95), y + Inches(0.04), Inches(3.9), Inches(0.5), t, 10.5, INK2,
         anchor=MSO_ANCHOR.MIDDLE)
box(s, Inches(8.4), Inches(5.4), Inches(4.43), Inches(1.25), PINK, PINK2)
text(s, Inches(8.6), Inches(5.5), Inches(4.1), Inches(1.1),
     [("配置のポイント", {"bold": True, "color": RED_D}),
      "・すべての什器・スタッフはライン内側",
      "・商談席から子どもが見える位置にキッズケア",
      "・混雑時はWS受付を整理券制に切替"], 9.5, INK, spacing=1.25)
footer(s, 8, "図はイメージ。実際の配置はロケハン後の設置図でご提出します。")

# ---------- 09 運営とKPI ----------
s = prs.slides.add_slide(BLANK)
header(s, "08", "運営フローとKPI ─ 「賑わい」で終わらせない",
       "各段階をプロセスKPIで管理し、2時間単位で未達のステップを特定。稼働中に打ち手を切り替えます。")
rows = [
    ["段階", "主役コンテンツ", "具体的な打ち手", "モニタリング指標"],
    ["視認・立寄", "エア遊具 / ゲーム遊具", "通路側に配置。装飾・POPで「何のイベントか」を一目で伝える", "通行者数・立寄数・立寄率"],
    ["参加", "ワークショップ", "参加条件として簡易アンケート（お得診断）を記入", "WS参加組数・参加率"],
    ["会話", "WS待ち時間", "制作中の保護者様へ通信・決済・固定回線をヒアリング", "接客化数・ヒアリング完了数"],
    ["商談", "イライラ棒（キッズケア）", "お子様を遊ばせつつ商談テーブルで落ち着いてご提案", "着座数・提案数・PI／副商材"],
]
table(s, Inches(0.5), Inches(1.6), Inches(12.33), rows,
      [Inches(1.6), Inches(2.4), Inches(5.13), Inches(3.2)], row_h=Inches(0.55),
      body_size=10, aligns=[C, L, L, L])
cards = [
    ("コンテンツ別の効果を記録",
     "会場ごとに「コンテンツによる集客・着火の規模」と「そこからの成約数」を毎日記録し、ライン編成の見直しに使います。"),
    ("二次策・三次策を事前準備",
     "主コンテンツの反応が弱い場合に備え、WSの入替（スクイーズ⇔チャーム）やアンケート起点への切替をロケハン時点で用意します。"),
    ("ヒアリングを飛ばさない",
     "WS待ち時間を活用し、ヒアリングを十分に行ってから提案に入る接客順序を徹底します。"),
]
for i, (h, b) in enumerate(cards):
    x = Inches(0.5) + Inches(4.17) * i
    box(s, x, Inches(4.75), Inches(3.95), Inches(2.1), PINK, PINK2)
    text(s, x + Inches(0.2), Inches(4.9), Inches(3.55), Inches(0.35), h, 12, RED_D, True)
    text(s, x + Inches(0.2), Inches(5.35), Inches(3.55), Inches(1.4), b, 10, INK2, spacing=1.3)
footer(s, 9)

# ---------- 10 スノードーム試験導入 ----------
s = prs.slides.add_slide(BLANK)
header(s, "09", "新ワークショップ「スノードーム」試験導入プラン",
       "現在試作段階のため、2ラインで小さく検証し、基準を満たした時点で横展開する段階導入とします。")
phases = [
    ("PHASE 1", "社内試作・安全確認", "〜11月上旬",
     ["材料・工程の確定（液漏れ・破損の確認）", "所要時間・1日あたり対応組数の計測", "材料原価と参加費設定の検討"]),
    ("PHASE 2", "2ラインで試験稼働", "11月中旬〜12月",
     ["L-07・L-08で冬季限定として稼働", "スクイーズ・チャームと参加率を比較", "保護者様の反応・着座への転換を記録"]),
    ("PHASE 3", "評価・横展開判断", "12月末",
     ["検証KPIを満たせば他ラインへ展開", "未達ならチャーム／スクイーズに戻す", "結果をご報告のうえ1月以降の編成に反映"]),
]
for i, (en, jp, when, its) in enumerate(phases):
    x = Inches(0.5) + Inches(4.17) * i
    y = Inches(1.65)
    hd = box(s, x, y, Inches(3.95), Inches(0.9), YEL if i != 1 else YEL_D,
             shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
    text(s, x + Inches(0.45), y + Inches(0.1), Inches(3.2), Inches(0.75),
         [(en + "　" + when, {"size": 9, "bold": True}), (jp, {"size": 14, "bold": True})], 12,
         WHITE if i == 1 else INK, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x, y + Inches(1.05), Inches(3.95), Inches(1.85), WHITE, LINE)
    text(s, x + Inches(0.2), y + Inches(1.2), Inches(3.55), Inches(1.6),
         ["▍" + t for t in its], 10, INK2, spacing=1.4)
rows = [
    ["検証KPI（案）", "基準（案）", "確認方法"],
    ["WS参加率（立寄に対する参加）", "スクイーズと同等以上", "時間帯別カウント"],
    ["1組あたり所要時間", "20分以内", "受付〜お渡しの実測"],
    ["保護者様の接客化率", "既存WSと同等以上", "日次実績シート"],
    ["安全・品質", "破損・液漏れ ゼロ", "お渡し時チェック"],
]
table(s, Inches(0.5), Inches(4.75), Inches(8.2), rows,
      [Inches(3.2), Inches(2.6), Inches(2.4)], row_h=Inches(0.38), body_size=9.5, aligns=[L, C, C])
photo(s, None, Inches(8.95), Inches(4.75), Inches(3.88), Inches(1.9))
footer(s, 10, "基準値は案です。試作結果を踏まえて確定します。")

# ---------- 11 施設ルール ----------
s = prs.slides.add_slide(BLANK)
header(s, "10", "施設ルールの順守 ─ 1Fメインで安心してご利用いただくために",
       "コンテンツ選定の段階から施設ルールを前提にし、「持ち込んだが使えない」を起こしません。")
rules = [
    ("リースライン内で完結",
     ["ロケハン時に施設様と実測し、ロケハンシートに記録", "初日朝に床面表示し全スタッフで共有",
      "ライン外での声掛けは行わない", "手を挙げて呼び込む等の目立つ所作は行わない"]),
    ("コンテンツの事前適合確認",
     ["面積・天井高・重量・電源・音出し可否を一覧化", "規制が厳しい施設は TYPE D（省スペース）を選択",
      "施設様の安全審査書類（保険証書・仕様書・設置図）を事前提出"]),
    ("店舗様・施設様との連携",
     ["事前MTGで受け入れキャパシティを数値で確認", "ご指摘は当日中に事務局長からご報告",
      "申し入れがあれば即時に活動停止し是正"]),
]
for i, (h, its) in enumerate(rules):
    x = Inches(0.5) + Inches(4.17) * i
    box(s, x, Inches(1.65), Inches(3.95), Inches(3.6), WHITE, LINE)
    circle_num(s, x + Inches(0.2), Inches(1.85), i + 1)
    text(s, x + Inches(0.75), Inches(1.88), Inches(3.0), Inches(0.4), h, 13, INK, True)
    text(s, x + Inches(0.2), Inches(2.5), Inches(3.55), Inches(2.7), ["▍" + t for t in its], 10,
         INK2, spacing=1.45)
callout(s, Inches(5.5), "直近の反省を反映",
        "直近の催事で「リースラインの順守」「施設条件に合わないコンテンツ」がご指摘となった事例を踏まえ、"
        "本パッケージでは型の選定段階で施設条件を確認し、ラインごとに順守状況をディレクターが1日3回巡回確認します。",
        Inches(1.1))
footer(s, 11)

# ---------- 12 季節展開 ----------
s = prs.slides.add_slide(BLANK)
header(s, "11", "季節ごとの展開イメージ（11月〜3月）",
       "同じ10ラインでも、時期に合わせてワークショップと装飾を入れ替え、リピーターにも新鮮さを保ちます。")
seasons = [
    ("11月", "秋〜ブラックフライデー", "通常編成で立ち上げ。スノードームはPHASE 2開始", RED),
    ("12月", "クリスマス", "TYPE C（スノードーム）を2ラインで本格稼働。装飾をクリスマス仕様に", YEL_D),
    ("1月", "お正月・初売り", "検証結果を反映し編成を再構築。連休はTYPE Aを優先配置", RED),
    ("2月", "冬休み明け・バレンタイン", "チャームをハート・ギフト仕様に。平日はTYPE B/Dで効率運営", BLUE_D),
    ("3月", "春休み・新生活", "家族来場が増える時期。TYPE A比率を高め集客を最大化", GREEN_D),
]
for i, (m, ev, act, col) in enumerate(seasons):
    x = Inches(0.5) + Inches(2.5) * i
    hd = box(s, x, Inches(1.7), Inches(2.33), Inches(0.75), col)
    label_in(hd, m, 18)
    box(s, x, Inches(2.5), Inches(2.33), Inches(2.6), WHITE, LINE)
    text(s, x + Inches(0.15), Inches(2.65), Inches(2.03), Inches(0.5), ev, 11, col, True)
    text(s, x + Inches(0.15), Inches(3.2), Inches(2.03), Inches(1.8), act, 10, INK2, spacing=1.35)
callout(s, Inches(5.45), "ねらい",
        "ワークショップの作品は「持ち帰って家族に見せる記念品」。季節テーマを変えることで、同じモールへの再来場と"
        "ドコモショップ様への来店予約につなげます。", Inches(0.9))
footer(s, 12, "季節施策は案です。各モールの販促カレンダーに合わせて調整します。")

# ---------- 13 次のステップ ----------
s = prs.slides.add_slide(BLANK)
header(s, "12", "次のステップとご確認事項",
       "本資料は構成案です。以下をすり合わせのうえ、正式なお見積りと稼働計画をご提出します。")
rows = [
    ["ご確認させていただきたい事項", "内容"],
    ["対象モール・稼働期間", "10ラインの展開候補モール、稼働開始時期・期間・曜日"],
    ["1Fメインの規格", "各モールの面積・天井高・電源・音出し可否・リースライン条件"],
    ["ライン構成のご希望", "TYPE A〜Dの比率、スノードーム試験導入の可否"],
    ["ご予算感", "1ラインあたりのご予算（コンテンツ費・運営人件費・送料の切り分け）"],
    ["運営体制", "スタッフ供給範囲（コンテンツ運営のみ／販売スタッフ込み）"],
]
table(s, Inches(0.5), Inches(1.6), Inches(7.6), rows, [Inches(2.5), Inches(5.1)],
      row_h=Inches(0.5), body_size=10, aligns=[L, L])
steps = [("STEP 1", "構成案のすり合わせ", "本資料をもとにご要望を伺う"),
         ("STEP 2", "ロケハン・型の確定", "候補モールの実測と TYPE 決定"),
         ("STEP 3", "お見積り・稼働計画", "ライン別の費用と稼働カレンダー"),
         ("STEP 4", "稼働開始", "日次KPI報告・週次の編成見直し")]
for i, (en, h, b) in enumerate(steps):
    y = Inches(1.6) + Inches(1.05) * i
    c = box(s, Inches(8.5), y, Inches(4.33), Inches(0.9), PINK if i else PINK2, PINK2)
    text(s, Inches(8.7), y + Inches(0.1), Inches(4.0), Inches(0.75),
         [(en, {"size": 9, "color": RED, "bold": True}), (h, {"size": 12, "bold": True}),
          (b, {"size": 9, "color": INK2})], 10, spacing=1.1)
box(s, Inches(0.5), Inches(5.95), Inches(12.33), Inches(0.9), RGBColor(0x3D, 0x2E, 0x3F))
text(s, Inches(0.8), Inches(6.02), Inches(11.8), Inches(0.75),
     [("お問い合わせ　株式会社DreamCreation", {"size": 12, "bold": True, "color": WHITE}),
      ("TEL 070-1319-2126　／　MAIL data@d-creation-o.com", {"size": 10, "color": WHITE})], 10,
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, 13)

prs.save(OUT)
print(OUT)
